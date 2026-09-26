from __future__ import annotations

import base64
from io import BytesIO

from docx import Document

from promptshield.schemas import ParsedInput, SourceType


def _to_bytes(content: bytes | str) -> bytes:
    if isinstance(content, bytes):
        return content
    return base64.b64decode(content)


def parse(content: bytes | str, source_type: SourceType, metadata: dict) -> ParsedInput:
    doc = Document(BytesIO(_to_bytes(content)))

    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells if cell.text)

    return ParsedInput(text="\n".join(parts), source_type=source_type, metadata=dict(metadata))
